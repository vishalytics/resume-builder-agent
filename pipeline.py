import hashlib
import json
import re
import tempfile
import unicodedata
import uuid
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document

from ats_checker import build_ats_report
from ats_optimizer import OPTIMIZER_CONFIG, optimize_cv
from cv_generator import (
    _validate_cv_response,
    apply_verified_source_facts,
    generate_cv,
    render_cv_document,
    validate_cv,
)
from gemini_client import (
    MAX_ATTEMPTS_PER_TASK,
    MODELS,
    RETRIES_PER_MODEL,
    redact_sensitive_text,
)
from interview_prep import build_interview_questions, render_interview_document
from jd_analyzer import _validate_jd_response, analyze_jd
from resume_matcher import _validate_match_response, match_resume_to_jd


BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
OUTPUT_DIR = BASE_DIR / "output"
MASTER_RESUME_PATH = BASE_DIR / "master_resume" / "master_resume.docx"
PIPELINE_CACHE_VERSION = "resume-agent-schema-4"
STAGE_NAMES = (
    "jd_analysis", "resume_matching", "cv_generation",
    "ats_optimization", "cv_document", "ats_report",
    "interview_preparation", "interview_document",
)


class PipelineError(RuntimeError):
    def __init__(self, run_id, stage, message, output_files=None, manifest_path=None):
        self.run_id = run_id
        self.stage = stage
        self.output_files = dict(output_files or {})
        self.manifest_path = Path(manifest_path) if manifest_path else None
        super().__init__(
            f"Run {run_id} failed during {stage}: {redact_sensitive_text(message)}"
        )


def create_run_id():
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    return f"{timestamp}_{uuid.uuid4().hex[:8]}"


def _sha256_bytes(value):
    return hashlib.sha256(value).hexdigest()


def _sha256_json(value):
    encoded = json.dumps(value, ensure_ascii=False, sort_keys=True).encode("utf-8")
    return _sha256_bytes(encoded)


def _sha256_file(path):
    path = Path(path)
    if not path.is_file():
        return ""
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _normalized_tokens(value):
    normalized = unicodedata.normalize("NFKD", str(value or ""))
    normalized = normalized.encode("ascii", "ignore").decode("ascii").lower()
    return set(re.findall(r"[a-z0-9]+", normalized))


def _date_tokens(value):
    tokens = _normalized_tokens(value)
    normalized = set()
    for token in tokens:
        compact_date = re.fullmatch(r"([a-z]{3,9})(\d{2})", token)
        if compact_date:
            normalized.add(compact_date.group(1))
            normalized.add(f"20{compact_date.group(2)}")
            continue
        if token.isdigit() and len(token) == 2:
            normalized.add(f"20{token}")
        else:
            normalized.add(token)
    return normalized


def _read_master_resume_text(path=MASTER_RESUME_PATH):
    document = Document(path)
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    return "\n".join(paragraphs)


def _resume_evidence_alignment_errors(resume_data, master_text):
    master_tokens = _normalized_tokens(master_text)
    errors = []
    candidate = resume_data.get("candidate", {})
    if not isinstance(candidate, dict):
        candidate = {}
    for field in ("name", "email", "phone", "linkedin", "location"):
        value = candidate.get(field, "")
        if field == "phone":
            supported = bool(re.sub(r"\D", "", str(value))) and re.sub(r"\D", "", str(value)) in re.sub(r"\D", "", master_text)
        elif field in ("email", "linkedin"):
            supported = bool(value) and str(value).casefold() in master_text.casefold()
        else:
            tokens = _normalized_tokens(value)
            supported = bool(tokens) and tokens <= master_tokens
        if value and not supported:
            errors.append(f"candidate.{field}")

    experiences = resume_data.get("experience", [])
    if not isinstance(experiences, list) or len(experiences) != 4:
        errors.append("experience (expected four documented roles)")
    else:
        for index, experience in enumerate(experiences, start=1):
            for field in ("company", "role"):
                tokens = _normalized_tokens(experience.get(field, ""))
                if not tokens or not tokens <= master_tokens:
                    errors.append(f"experience[{index}].{field}")
            dates = experience.get("duration", experience.get("dates", ""))
            date_tokens = _date_tokens(dates)
            if dates and (not date_tokens or not date_tokens <= master_tokens):
                errors.append(f"experience[{index}].duration")

    for index, education in enumerate(resume_data.get("education", []), start=1):
        for field in ("degree", "institution", "duration"):
            tokens = _normalized_tokens(education.get(field, ""))
            if not tokens or not tokens <= master_tokens:
                errors.append(f"education[{index}].{field}")
        score = education.get("cgpa", education.get("score", ""))
        if score and _normalized_tokens(score) - master_tokens:
            errors.append(f"education[{index}].score")

    for index, certification in enumerate(resume_data.get("certifications", []), start=1):
        title = certification.get("name", "") if isinstance(certification, dict) else certification
        tokens = _normalized_tokens(title)
        if not tokens or not tokens <= master_tokens:
            errors.append(f"certifications[{index}]")
    return errors


def _master_resume_highlights(master_text):
    achievements = []
    leadership = []
    section = None
    for raw_line in master_text.splitlines():
        line = raw_line.strip().lstrip("•●▪-* ").strip()
        upper = line.upper()
        if "SELECTED ACHIEVEMENTS" in upper:
            section = "achievements"
            continue
        if upper.startswith("CERTIFICATIONS"):
            section = None
            continue
        if upper == "LEADERSHIP":
            section = "leadership"
            continue
        if not line or section is None:
            continue
        if section == "achievements" and line.startswith("Academic Manager"):
            section = "leadership"
        if section == "achievements":
            clean_line = re.sub(
                r"\s*\((?:state that|confirm|validate|add your|reconcile|use this wording|verify exact).*?\)",
                "", line, flags=re.IGNORECASE,
            ).strip()
            if clean_line.startswith("Received a full "):
                clean_line = re.split(r"\s+after selection\b", clean_line, maxsplit=1, flags=re.IGNORECASE)[0].rstrip(" ,;") + "."
            if clean_line:
                achievements.append(clean_line)
        if section == "leadership":
            if any(marker in line.casefold() for marker in (
                "add your exact role", "confirm your role", "validate exact",
            )):
                continue
            leadership.extend(part.strip() for part in line.split(";") if part.strip())
    return achievements, leadership


def _load_json(path):
    with Path(path).open("r", encoding="utf-8") as file:
        value = json.load(file)
    if not isinstance(value, dict):
        raise ValueError(f"Expected a JSON object in {path}.")
    return value


def _write_json(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = path.with_suffix(path.suffix + ".tmp")
    with temporary_path.open("w", encoding="utf-8") as file:
        json.dump(value, file, indent=2, ensure_ascii=False)
    temporary_path.replace(path)


def _slug(value, fallback):
    return re.sub(r"[^A-Za-z0-9_-]+", "_", str(value or "")).strip("_") or fallback


def _path_reference(path):
    path = Path(path).resolve()
    try:
        return str(path.relative_to(BASE_DIR))
    except ValueError:
        return str(path)


def _valid_jd(value):
    return isinstance(value, dict) and _validate_jd_response(json.dumps(value))[0]


def _valid_match(value):
    return isinstance(value, dict) and _validate_match_response(json.dumps(value))[0]


def _valid_cv(value):
    return isinstance(value, dict) and validate_cv(value)[0]


def _valid_optimization(value):
    selected_report = value.get("selected_report", {}) if isinstance(value, dict) else {}
    optimization = value.get("optimization", {}) if isinstance(value, dict) else {}
    document_integrity = selected_report.get("document_integrity", {}) if isinstance(selected_report, dict) else {}
    return (
        isinstance(value, dict)
        and _valid_cv(value.get("best_cv"))
        and isinstance(value.get("selected_report"), dict)
        and isinstance(value.get("optimization"), dict)
        and isinstance(value["selected_report"].get("ats_compatibility_score"), (int, float))
        and isinstance(value["optimization"].get("history"), list)
        and document_integrity.get("checked") is True
        and document_integrity.get("valid") is True
        and document_integrity.get("source") == "rendered_docx"
        and optimization.get("final_score") == selected_report.get("ats_compatibility_score")
    )


def _read_docx_text(path):
    document = Document(path)
    sections = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            sections.extend(cell.text for cell in row.cells)
    return "\n".join(text for text in sections if text)


def _valid_document_file(path, stage):
    try:
        with zipfile.ZipFile(path) as package:
            if package.testzip() is not None:
                return False
        document = Document(path)
    except (OSError, ValueError, zipfile.BadZipFile):
        return False
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    if stage == "cv_document":
        return (
            "EDUCATION" in text
            and "PROFESSIONAL EXPERIENCE" in text
            and "CERTIFICATIONS" in text
            and "{{" not in text
            and "}}" not in text
        )
    if stage == "interview_document":
        return "Interview Preparation" in text and "Genuine experience gaps" in text
    return False


def _valid_interview(value):
    questions = value.get("questions") if isinstance(value, dict) else None
    gaps = value.get("experience_gaps") if isinstance(value, dict) else None
    return (
        isinstance(questions, list)
        and 4 <= len(questions) <= 10
        and any(item.get("category") == "analytical_case" for item in questions if isinstance(item, dict))
        and isinstance(gaps, list)
        and all(
            isinstance(item, dict)
            and isinstance(item.get("question"), str)
            and item.get("category") in {"role_specific", "behavioral", "analytical_case"}
            and isinstance(item.get("answer_points"), list)
            and isinstance(item.get("follow_up_questions"), list)
            and isinstance(item.get("evidence_indices"), list)
            and all(type(index) is int and index >= 0 for index in item.get("evidence_indices", []))
            for item in questions
        )
    )


def _manifest_write(manifest_path, manifest):
    manifest["updated_at"] = datetime.now(timezone.utc).isoformat(timespec="seconds")
    _write_json(manifest_path, manifest)


def _callback(callback, stage, status, detail=""):
    if callback is not None:
        callback(stage, status, detail)


def run_pipeline(jd_text, *, run_id=None, progress_callback=None):
    if not isinstance(jd_text, str) or not jd_text.strip():
        raise ValueError("Enter a job description before generating an application pack.")
    run_id = run_id or create_run_id()
    if not re.fullmatch(r"[A-Za-z0-9_-]{1,80}", run_id):
        raise ValueError("Run ID contains invalid path characters.")

    evidence_path = DATA_DIR / "resume_evidence.json"
    resume_data = _load_json(evidence_path)
    master_resume_text = _read_master_resume_text()
    alignment_errors = _resume_evidence_alignment_errors(resume_data, master_resume_text)
    if alignment_errors:
        raise ValueError(
            "Stored resume evidence does not align with the unchanged master resume: "
            + ", ".join(alignment_errors)
            + ". Refresh the evidence database from the master before applying."
        )
    master_resume_hash = _sha256_file(MASTER_RESUME_PATH)
    resume_hash = _sha256_json(resume_data)
    jd_hash = _sha256_bytes(jd_text.strip().encode("utf-8"))
    cache_configuration = {
        "version": PIPELINE_CACHE_VERSION,
        "models": MODELS,
        "retries_per_model": RETRIES_PER_MODEL,
        "max_attempts_per_task": MAX_ATTEMPTS_PER_TASK,
        "optimizer": OPTIMIZER_CONFIG,
    }
    base_fingerprint = _sha256_json({
        "jd": jd_hash,
        "resume_evidence": resume_hash,
        "master_resume": master_resume_hash,
        "configuration": cache_configuration,
    })

    data_run_dir = DATA_DIR / "runs" / run_id
    output_run_dir = OUTPUT_DIR / "runs" / run_id
    data_run_dir.mkdir(parents=True, exist_ok=True)
    output_run_dir.mkdir(parents=True, exist_ok=True)
    manifest_path = data_run_dir / "run_manifest.json"
    jd_text_path = data_run_dir / "job_description.txt"
    if manifest_path.exists():
        manifest = _load_json(manifest_path)
        if manifest.get("base_fingerprint") != base_fingerprint:
            raise PipelineError(
                run_id, "input_validation",
                "This run ID belongs to different JD or resume inputs; create a new run.",
                manifest.get("outputs", {}), manifest_path,
            )
    else:
        jd_text_path.write_text(jd_text, encoding="utf-8")
        _write_json(data_run_dir / "resume_evidence.json", resume_data)
        manifest = {
            "run_id": run_id,
            "status": "running",
            "created_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
            "base_fingerprint": base_fingerprint,
            "input": {
                "job_description_sha256": jd_hash,
                "resume_evidence_sha256": resume_hash,
                "master_resume_sha256": master_resume_hash,
            },
            "stages": {},
            "outputs": {},
            "model_attempts": [],
        }
        _manifest_write(manifest_path, manifest)

    manifest["status"] = "running"
    manifest.pop("error", None)
    _manifest_write(manifest_path, manifest)
    active_stage = "initialization"

    def output_files():
        return {
            name: str(BASE_DIR / relative_path)
            for name, relative_path in manifest.get("outputs", {}).items()
            if (BASE_DIR / relative_path).is_file()
        }

    def cached_stage(stage, input_fingerprint, json_path, validator):
        entry = manifest.get("stages", {}).get(stage, {})
        if entry.get("status") not in ("fresh", "reused"):
            return None
        if entry.get("input_fingerprint") != input_fingerprint or not Path(json_path).is_file():
            return None
        try:
            value = _load_json(json_path)
        except (OSError, ValueError, json.JSONDecodeError):
            return None
        if not validator(value):
            return None
        entry["status"] = "reused"
        entry["reused_at"] = datetime.now(timezone.utc).isoformat(timespec="seconds")
        _manifest_write(manifest_path, manifest)
        _callback(progress_callback, stage, "reused", "Validated output reused for this run.")
        return value

    def run_json_stage(stage, fingerprint, json_path, compute, validator):
        nonlocal active_stage
        active_stage = stage
        cached = cached_stage(stage, fingerprint, json_path, validator)
        if cached is not None:
            return cached
        events = []
        manifest.setdefault("stages", {})[stage] = {
            "status": "running",
            "input_fingerprint": fingerprint,
            "started_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
            "model_attempts": [],
        }
        _manifest_write(manifest_path, manifest)
        _callback(progress_callback, stage, "running")
        try:
            value = compute(events)
            if not validator(value):
                raise ValueError(f"Stage {stage} returned invalid structured data.")
            _write_json(json_path, value)
            manifest["stages"][stage].update({
                "status": "fresh",
                "completed_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
                "model_attempts": events,
                "artifact": _path_reference(json_path),
            })
            manifest.setdefault("model_attempts", []).extend(events)
            _manifest_write(manifest_path, manifest)
            _callback(progress_callback, stage, "fresh")
            return value
        except Exception as error:
            events.extend(event for event in getattr(error, "events", ()) if event not in events)
            manifest["stages"][stage].update({
                "status": "failed",
                "completed_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
                "error": redact_sensitive_text(error),
                "model_attempts": events,
            })
            manifest.setdefault("model_attempts", []).extend(
                event for event in events if event not in manifest["model_attempts"]
            )
            manifest["status"] = "failed"
            manifest["failed_stage"] = stage
            manifest["error"] = redact_sensitive_text(error)
            _manifest_write(manifest_path, manifest)
            _callback(progress_callback, stage, "failed", redact_sensitive_text(error))
            raise PipelineError(run_id, stage, str(error), output_files(), manifest_path) from None

    def run_file_stage(stage, fingerprint, destination, writer):
        nonlocal active_stage
        active_stage = stage
        entry = manifest.get("stages", {}).get(stage, {})
        if (
            entry.get("status") in ("fresh", "reused")
            and entry.get("input_fingerprint") == fingerprint
            and Path(destination).is_file()
            and _valid_document_file(destination, stage)
        ):
            entry["status"] = "reused"
            _manifest_write(manifest_path, manifest)
            _callback(progress_callback, stage, "reused", "Validated file reused for this run.")
            return Path(destination)
        manifest.setdefault("stages", {})[stage] = {
            "status": "running",
            "input_fingerprint": fingerprint,
            "started_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        }
        _manifest_write(manifest_path, manifest)
        _callback(progress_callback, stage, "running")
        try:
            result_path = Path(writer())
            if not result_path.is_file() or result_path.stat().st_size == 0:
                raise ValueError(f"Stage {stage} did not create a non-empty output file.")
            if not _valid_document_file(result_path, stage):
                raise ValueError(f"Stage {stage} produced an invalid or incomplete DOCX file.")
            manifest["stages"][stage].update({
                "status": "fresh",
                "completed_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
                "artifact": _path_reference(result_path),
            })
            output_name = "cv" if stage == "cv_document" else "interview" if stage == "interview_document" else "ats_report"
            manifest.setdefault("outputs", {})[output_name] = _path_reference(result_path)
            _manifest_write(manifest_path, manifest)
            _callback(progress_callback, stage, "fresh")
            return result_path
        except Exception as error:
            manifest["stages"][stage].update({
                "status": "failed",
                "completed_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
                "error": redact_sensitive_text(error),
            })
            manifest["status"] = "failed"
            manifest["failed_stage"] = stage
            manifest["error"] = redact_sensitive_text(error)
            _manifest_write(manifest_path, manifest)
            _callback(progress_callback, stage, "failed", redact_sensitive_text(error))
            raise PipelineError(run_id, stage, str(error), output_files(), manifest_path) from None

    try:
        jd_fingerprint = _sha256_json((base_fingerprint, "jd_analysis"))
        jd_data = run_json_stage(
            "jd_analysis", jd_fingerprint, data_run_dir / "jd_analysis.json",
            lambda events: analyze_jd(jd_text, event_sink=events), _valid_jd,
        )

        match_fingerprint = _sha256_json((base_fingerprint, "resume_matching", jd_data))
        match_data = run_json_stage(
            "resume_matching", match_fingerprint, data_run_dir / "resume_jd_match.json",
            lambda events: match_resume_to_jd(resume_data, jd_data, event_sink=events), _valid_match,
        )

        cv_fingerprint = _sha256_json((base_fingerprint, "cv_generation", jd_data, match_data))
        cv_data = run_json_stage(
            "cv_generation", cv_fingerprint, data_run_dir / "generated_cv.json",
            lambda events: generate_cv(resume_data, jd_data, match_data, event_sink=events), _valid_cv,
        )
        cv_data = apply_verified_source_facts(cv_data, resume_data, jd_data)
        master_achievements, master_leadership = _master_resume_highlights(master_resume_text)
        if master_achievements:
            cv_data["achievements"] = master_achievements
        if master_leadership:
            cv_data["leadership"] = ["; ".join(master_leadership)]
        valid_cv, validation_message = validate_cv(cv_data)
        if not valid_cv:
            raise PipelineError(run_id, "cv_validation", validation_message, output_files(), manifest_path)

        def evaluate_cv_candidate(candidate_cv):
            valid, message = validate_cv(candidate_cv)
            if not valid:
                raise ValueError(f"Optimization candidate failed CV validation: {message}")
            with tempfile.TemporaryDirectory(prefix="resume-agent-cv-eval-") as temporary_directory:
                candidate_path = render_cv_document(candidate_cv, jd_data, temporary_directory)
                if not _valid_document_file(candidate_path, "cv_document"):
                    raise ValueError("Optimization candidate produced an invalid CV document.")
                document_text = _read_docx_text(candidate_path)
                report = build_ats_report(
                    jd_data,
                    match_data,
                    candidate_cv,
                    resume_data,
                    document_text=document_text,
                    document_integrity={"checked": True, "valid": True, "source": "rendered_docx"},
                )
                return report, document_text

        optimization_fingerprint = _sha256_json((
            base_fingerprint, "ats_optimization", cv_fingerprint,
            jd_data, match_data, OPTIMIZER_CONFIG,
        ))

        def run_optimization(_events):
            best_cv, selected_report, optimization_result = optimize_cv(
                cv_data,
                jd_data,
                match_data,
                resume_data,
                evaluate_cv_candidate,
                config=OPTIMIZER_CONFIG,
                master_text=master_resume_text,
            )
            return {
                "best_cv": best_cv,
                "selected_report": selected_report,
                "optimization": optimization_result,
            }

        optimization_record = run_json_stage(
            "ats_optimization",
            optimization_fingerprint,
            data_run_dir / "ats_optimization.json",
            run_optimization,
            _valid_optimization,
        )
        cv_data = optimization_record["best_cv"]
        optimization_result = optimization_record["optimization"]
        _write_json(data_run_dir / "generated_cv.json", cv_data)

        candidate_slug = _slug(cv_data.get("name"), "Candidate")
        role_slug = _slug(jd_data.get("job_title"), "Target_Role")
        cv_path = output_run_dir / f"{candidate_slug}_{role_slug}_CV.docx"
        cv_doc_fingerprint = _sha256_json((optimization_fingerprint, "cv_document", cv_data))
        cv_path = run_file_stage(
            "cv_document", cv_doc_fingerprint, cv_path,
            lambda: render_cv_document(cv_data, jd_data, output_run_dir),
        )

        selected_document_text = _read_docx_text(cv_path)
        report_fingerprint = _sha256_json((
            base_fingerprint, "ats_report", jd_data, match_data, cv_data,
            _sha256_file(cv_path), optimization_result,
        ))
        report_path = output_run_dir / "ats_report.json"

        def build_selected_report(_events):
            report = build_ats_report(
                jd_data,
                match_data,
                cv_data,
                resume_data,
                document_text=selected_document_text,
                document_integrity={"checked": True, "valid": _valid_document_file(cv_path, "cv_document"), "source": "selected_docx"},
            )
            report["optimization"] = optimization_result
            report["optimization"]["score_from_selected_docx"] = True
            report["optimization"]["selected_docx_score"] = report["ats_compatibility_score"]
            report["run_id"] = run_id
            return report

        report = run_json_stage(
            "ats_report", report_fingerprint, report_path,
            build_selected_report,
            lambda value: isinstance(value.get("ats_compatibility_score"), (int, float))
            and value.get("run_id") == run_id,
        )
        if report["ats_compatibility_score"] != optimization_result["final_score"]:
            raise ValueError("Final ATS score does not match the selected rendered-CV score.")
        manifest.setdefault("outputs", {})["ats_report"] = _path_reference(report_path)
        _manifest_write(manifest_path, manifest)

        interview_fingerprint = _sha256_json((base_fingerprint, "interview_preparation", jd_data, match_data))
        interview_content = run_json_stage(
            "interview_preparation", interview_fingerprint,
            data_run_dir / "interview_preparation.json",
            lambda events: build_interview_questions(resume_data, jd_data, match_data, event_sink=events),
            _valid_interview,
        )
        interview_path = output_run_dir / f"{candidate_slug}_{role_slug}_Interview_Preparation.docx"
        interview_doc_fingerprint = _sha256_json((interview_fingerprint, "interview_document"))
        interview_path = run_file_stage(
            "interview_document", interview_doc_fingerprint, interview_path,
            lambda: render_interview_document(interview_content, jd_data, interview_path),
        )
        for name, path in (("cv", cv_path), ("interview", interview_path), ("ats_report", report_path)):
            manifest.setdefault("outputs", {})[name] = _path_reference(path)
        manifest["status"] = "completed"
        manifest["completed_at"] = datetime.now(timezone.utc).isoformat(timespec="seconds")
        manifest.pop("failed_stage", None)
        _manifest_write(manifest_path, manifest)
        _callback(progress_callback, "pipeline", "completed")
        return {
            "run_id": run_id,
            "status": "completed",
            "stages": {name: manifest["stages"][name]["status"] for name in STAGE_NAMES},
            "jd_analysis": jd_data,
            "resume_match": match_data,
            "generated_cv": cv_data,
            "ats_report": report,
            "optimization": optimization_result,
            "files": {"cv": cv_path, "interview": interview_path, "ats_report": report_path},
            "manifest": manifest_path,
        }
    except PipelineError:
        raise
    except Exception as error:
        manifest["status"] = "failed"
        manifest["failed_stage"] = active_stage
        manifest["error"] = redact_sensitive_text(error)
        _manifest_write(manifest_path, manifest)
        _callback(progress_callback, active_stage, "failed", redact_sensitive_text(error))
        raise PipelineError(run_id, active_stage, str(error), output_files(), manifest_path) from None