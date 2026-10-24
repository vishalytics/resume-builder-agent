import json
import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from gemini_client import generate_with_fallback, parse_json_response


# ============================================================
# CONFIG
# ============================================================

BASE_DIR = Path(__file__).resolve().parent
RESUME_FILE = BASE_DIR / "data" / "resume_evidence.json"
JD_FILE = BASE_DIR / "data" / "jd_analysis.json"
MATCH_FILE = BASE_DIR / "data" / "resume_jd_match.json"

# This template is based on the optimized master resume layout.
TEMPLATE_FILE = BASE_DIR / "templates" / "CV_TEMPLATE.docx"

OUTPUT_DIR = BASE_DIR / "output"
GENERATED_JSON_FILE = BASE_DIR / "data" / "generated_cv.json"
EXPERIENCE_BULLET_LIMITS = {1: 8, 2: 8, 3: 5, 4: 2}

ACTION_VERBS = frozenset({
    "automated", "built", "conducted", "coordinated", "created", "delivered",
    "developed", "directed", "engineered", "evaluated", "executed", "generated",
    "improved", "implemented", "increased", "identified", "launched", "led",
    "managed", "onboarded", "optimized", "optimised", "orchestrated", "oversaw",
    "partnered", "reduced", "redesigned", "secured", "standardized", "streamlined",
    "transformed", "converted", "mentored", "negotiated", "supported",
})

ACHIEVEMENT_PHRASES = (
    "cross-functional rollout", "cross-functional teams", "real-time dashboards",
    "vendor-slot booking platform", "Product Master Cost Sheet", "Power BI dashboards",
    "supplier performance", "process compliance", "operational efficiency",
    "pricing intelligence", "stakeholder coordination",
)

MONEY_PATTERN = re.compile(
    r"(?:₹|INR\s*|Rs\.?\s*|USD\s*\$?\s*|\$)\s*"
    r"\d[\d,]*(?:\.\d+)?(?:\s?(?:lakhs?|crores?|million|billion|Cr|L|K|M))?\+?",
    re.IGNORECASE,
)
PERCENT_PATTERN = re.compile(r"\b\d[\d,]*(?:\.\d+)?\s?%")
SCALE_PATTERN = re.compile(
    r"\b\d[\d,]*(?:\.\d+)?(?:\+)?"
    r"(?:\s?-\s?(?:member|person|unit))?"
    r"\s+(?:(?:residential|academic|regional)\s+)?(?:students?|vendors?|members?|teams?|employees?|centres?|centers?|"
    r"campuses?|units?|clients?|orders?|SOPs?|questions?|products?|URLs?|buyers?|"
    r"scholars?|leads?|accounts?|projects?|years?|people)\b",
    re.IGNORECASE,
)


def load_json(path):
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def compact_resume_data(data):
    """Keep source evidence needed for a tailored CV, including nested identity."""
    keys = [
        "candidate", "name", "email", "phone", "linkedin", "location",
        "summary", "education", "experience", "work_experience",
        "professional_experience", "employment", "skills", "technical_skills",
        "tools_and_technologies", "achievements", "certifications",
        "projects", "leadership",
    ]
    return {key: data[key] for key in keys if key in data}


def compact_jd_data(data):
    keys = [
        "job_title", "company", "responsibilities", "required_skills",
        "preferred_skills", "tools_and_technologies", "domain_keywords", "ats_keywords",
        "behavioral_requirements", "experience_requirements",
        "education_requirements", "important_action_verbs",
        "key_business_outcomes",
    ]
    return {key: data[key] for key in keys if key in data}


def compact_match_data(data):
    keys = [
        "overall_match", "strong_matches", "partial_matches", "no_evidence",
        "keyword_matches", "missing_keywords", "skill_gaps", "priority_changes",
    ]
    return {key: data[key] for key in keys if key in data}


def validate_cv(data):
    if not isinstance(data, dict):
        return False, "Gemini response is not a JSON object."

    required = [
        "name", "targeted_summary", "core_expertise",
        "tools_and_technologies", "experience", "education",
        "achievements", "certifications",
    ]
    for field in required:
        if field not in data:
            return False, f"Missing field: {field}"

    if not str(data.get("name", "")).strip():
        return False, "Name is empty."
    if not str(data.get("targeted_summary", "")).strip():
        return False, "Targeted summary is empty."
    if not isinstance(data.get("experience"), list):
        return False, "Experience must be a list."
    if not isinstance(data.get("education"), dict):
        return False, "Education must be an object."
    for field in ("core_expertise", "tools_and_technologies", "achievements", "certifications"):
        if not isinstance(data.get(field), list):
            return False, f"{field} must be a list."

    for i, exp in enumerate(data["experience"], start=1):
        if not isinstance(exp, dict):
            return False, f"Experience entry {i} is not an object."
        if not isinstance(exp.get("bullets", []), list):
            return False, f"Bullets for experience entry {i} must be a list."

    return True, "Valid CV JSON."


def _validate_cv_response(response_text):
    try:
        return validate_cv(parse_json_response(response_text))
    except ValueError:
        return False, "Gemini response did not contain a valid CV object."



def normalize_source_certifications(value):
    """Normalize certification entries from resume_evidence.json, never from model output."""
    if value is None:
        return []

    raw_items = value if isinstance(value, list) else [value]
    entries = []
    for item in raw_items:
        if isinstance(item, dict):
            item = (item.get("name") or item.get("certification") or
                    item.get("title") or item.get("credential") or "")
        if not isinstance(item, str):
            continue
        # The parser may preserve the whole CERTIFICATIONS line as one string.
        for part in re.split(r"\s*[•●▪]\s*|\r?\n+", item):
            part = re.sub(r"^\s*[-*]+\s*", "", part).strip()
            if not part:
                continue
            # Guard against the exact failure seen: education copied into certifications.
            education_markers = (
                "mba degree", "b.tech degree", "btech degree",
                "indian institute of technology", "national institute of technology",
                "cgpa:", "2023–2025", "2023-2025", "2014–2018", "2014-2018"
            )
            if any(marker in part.lower() for marker in education_markers):
                continue
            if part not in entries:
                entries.append(part)
    return entries


def normalize_source_education(value):
    normalized = {
        "mba_degree": "", "mba_institution": "", "mba_duration": "", "mba_score": "",
        "btech_degree": "", "btech_institution": "", "btech_duration": "", "btech_score": "",
    }
    entries = value if isinstance(value, list) else [value]
    for entry in entries:
        if not isinstance(entry, dict):
            continue
        degree = str(entry.get("degree", "") or "").strip()
        degree_key = degree.lower().replace(".", "")
        prefix = "mba" if "mba" in degree_key else "btech" if "btech" in degree_key else ""
        if not prefix:
            continue
        normalized[f"{prefix}_degree"] = degree
        normalized[f"{prefix}_institution"] = str(entry.get("institution", "") or "").strip()
        normalized[f"{prefix}_duration"] = str(entry.get("duration", "") or "").strip()
        normalized[f"{prefix}_score"] = str(entry.get("cgpa", entry.get("score", "")) or "").strip()
    return normalized


def normalize_source_entries(value, fields):
    if value is None:
        return []
    items = value if isinstance(value, list) else [value]
    entries = []
    for item in items:
        if isinstance(item, dict):
            parts = [str(item.get(field, "") or "").strip() for field in fields]
            item = " - ".join(part for part in parts if part)
        if isinstance(item, str) and item.strip() and item.strip() not in entries:
            entries.append(item.strip())
    return entries


def apply_verified_source_facts(cv_data, resume_data, jd_data):
    candidate = resume_data.get("candidate", {})
    if not isinstance(candidate, dict):
        candidate = {}

    aliases = {
        "name": ("name", "full_name", "candidate_name"),
        "email": ("email", "email_address"),
        "phone": ("phone", "mobile", "phone_number"),
        "linkedin": ("linkedin", "linkedin_url", "linkedin_profile"),
        "location": ("location", "city", "current_location"),
    }
    for field, keys in aliases.items():
        source_value = next(
            (container[key] for container in (resume_data, candidate)
             for key in keys if container.get(key)),
            None,
        )
        if source_value:
            cv_data[field] = str(source_value)

    source_education = normalize_source_education(resume_data.get("education", []))
    if any(source_education.values()):
        cv_data["education"] = source_education

    source_experience = resume_data.get("experience", [])
    if isinstance(source_experience, list):
        model_experience = cv_data.get("experience", [])
        if not isinstance(model_experience, list):
            model_experience = []
        validated_experience = []
        for index, source in enumerate(source_experience[:4]):
            if not isinstance(source, dict):
                continue
            entry = model_experience[index] if index < len(model_experience) else {}
            if not isinstance(entry, dict):
                entry = {}
            for target, candidates in (
                ("company", ("company",)),
                ("role", ("role", "title")),
                ("dates", ("duration", "dates", "date")),
            ):
                source_value = next((source[key] for key in candidates if source.get(key)), None)
                if source_value:
                    entry[target] = str(source_value)
            source_text = json.dumps(source, ensure_ascii=False)
            source_numbers = set(re.findall(r"\d[\d,]*(?:\.\d+)?(?:\+|%)?", source_text))
            bullets = [
                bullet for bullet in as_list(entry.get("bullets", []))
                if all(
                    number in source_numbers
                    for number in re.findall(r"\d[\d,]*(?:\.\d+)?(?:\+|%)?", bullet)
                )
            ]
            if not bullets:
                for achievement in source.get("achievements", []):
                    if isinstance(achievement, dict):
                        text = achievement.get("evidence") or achievement.get("achievement")
                    else:
                        text = achievement
                    if isinstance(text, str) and text.strip():
                        bullets.append(text.strip())
                    if len(bullets) >= len(source.get("achievements", [])):
                        break
            entry["bullets"] = bullets
            validated_experience.append(entry)
        cv_data["experience"] = validated_experience

    cv_data["achievements"] = normalize_source_entries(
        resume_data.get("achievements", []), ("name", "achievement", "description")
    )
    cv_data["leadership"] = normalize_source_entries(
        resume_data.get("leadership", []), ("role", "institution", "duration")
    )
    cv_data["certifications"] = normalize_source_certifications(
        resume_data.get("certifications", [])
    )
    source_text = json.dumps(resume_data, ensure_ascii=False)
    source_numbers = set(re.findall(r"\d[\d,]*(?:\.\d+)?(?:\+|%)?", source_text))
    summary = str(cv_data.get("targeted_summary", "") or "")
    summary_numbers = re.findall(r"\d[\d,]*(?:\.\d+)?(?:\+|%)?", summary)
    if any(number not in source_numbers for number in summary_numbers):
        verified_summary = resume_data.get("summary", "")
        if isinstance(verified_summary, str) and verified_summary.strip():
            cv_data["targeted_summary"] = verified_summary.strip()
    evidence_text = json.dumps(resume_data, ensure_ascii=False).lower()
    evidence_tokens = set(re.findall(r"[a-z0-9+#.]+", evidence_text))

    def is_evidence_supported(value):
        tokens = set(re.findall(r"[a-z0-9+#.]+", str(value or "").lower()))
        return bool(tokens) and tokens <= evidence_tokens

    cv_data["core_expertise"] = [
        item for item in as_list(cv_data.get("core_expertise", []))
        if is_evidence_supported(item)
    ]
    verified_tools = []
    for line in as_list(cv_data.get("tools_and_technologies", [])):
        if ":" in line:
            category, values = line.split(":", 1)
            tools = [item.strip() for item in re.split(r"[,|\n]+", values) if item.strip()]
            supported_tools = [item for item in tools if is_evidence_supported(item)]
            if supported_tools:
                verified_tools.append(f"{category.strip()}: {', '.join(supported_tools)}")
        elif is_evidence_supported(line):
            verified_tools.append(line)
    cv_data["tools_and_technologies"] = verified_tools
    cv_data["targeted_headline"] = str(
        cv_data.get("targeted_headline") or jd_data.get("job_title", "") or ""
    )
    return cv_data

def generate_cv(resume_data, jd_data, match_data, event_sink=None):
    resume = compact_resume_data(resume_data)
    jd = compact_jd_data(jd_data)
    match = compact_match_data(match_data)

    prompt = f"""
You are an expert ATS resume editor. Produce CV CONTENT ONLY; the Word template controls
the layout, typography, heading names, and section order.

TARGET ROLE:
{jd.get("job_title", "Target Role")}

NON-NEGOTIABLE RULES:
- Use only facts supported by CANDIDATE EVIDENCE below. Do not invent employers, titles,
  dates, qualifications, tools, metrics, achievements, or certifications.
- Keep the experience entries in the same order as the source resume.
- Tailor wording to the target role only where the source evidence supports it.
- Use relevant JD terminology when supported by the candidate evidence, including analytical problem-solving, data analysis, stakeholder management, and cross-functional collaboration.
- Describe demonstrated analytical work through concrete evidence such as procurement analytics, performance dashboards, process improvement, and data-driven decisions; do not claim unsupported skills.
- Prefer clear, natural phrases such as "worked with cross-functional teams" when the evidence shows collaboration across departments.
- Never include editorial notes such as "confirm", "validate", "add your exact role",
  "reported notes cite", or similar instructions in the final CV.
- Do not add or rename resume section headings. The Word template supplies the headings:
  EDUCATION; HOW I CREATE BUSINESS IMPACT; CORE COMPETENCIES; TOOLS & TECHNOLOGY;
  PROFESSIONAL EXPERIENCE; SELECTED ACHIEVEMENTS & LEADERSHIP; CERTIFICATIONS.
- Keep the summary concise (2–3 lines). Prefer measurable results when verified.
- Return valid JSON only, with no markdown fences or explanations.
- Use empty strings/lists for missing information. Never output placeholder braces.

JOB REQUIREMENTS:
{json.dumps(jd, ensure_ascii=False)}

RESUME-JD MATCH:
{json.dumps(match, ensure_ascii=False)}

CANDIDATE EVIDENCE:
{json.dumps(resume, ensure_ascii=False)}

Return exactly this JSON structure:
{{
  "name": "",
  "email": "",
  "phone": "",
  "linkedin": "",
  "location": "",
  "targeted_headline": "",
  "targeted_summary": "",
  "core_expertise": [],
  "tools_and_technologies": [],
  "experience": [
    {{
      "company": "",
      "role": "",
      "dates": "",
      "focus": "",
      "bullets": []
    }}
  ],
  "achievements": [],
  "leadership": [],
  "certifications": [],
  "education": {{
    "mba_degree": "",
    "mba_institution": "",
    "mba_duration": "",
    "mba_score": "",
    "btech_degree": "",
    "btech_institution": "",
    "btech_duration": "",
    "btech_score": ""
  }}
}}

Additional requirements:
- Return every source experience entry exactly once, in source order.
- Use up to 8 concise bullets for the first two roles, up to 5 for the third, and up to 2 for the fourth.
- "focus" should be a short line of verified themes, without inventing skills.
- "core_expertise" should contain concise competencies supported by the evidence.
- "tools_and_technologies" should be a list of short category lines, such as
  "Analytics: Excel, Power BI, SQL" and "Automation & AI: Google Apps Script, Gemini APIs".
- Put leadership roles in "leadership"; do not duplicate them in "achievements".
- Preserve degree, institution, dates, and scores as written in the evidence.
- Do not include projects as a separate section because the master layout has no projects heading.
"""


    print("\nSending request to Gemini...")
    result = generate_with_fallback(
        prompt,
        stage="cv_generation",
        response_validator=_validate_cv_response,
        event_sink=event_sink,
    )
    print("Gemini response received. Response length:", len(result))

    data = parse_json_response(result)
    apply_verified_source_facts(data, resume_data, jd)

    valid, message = validate_cv(data)
    if not valid:
        print("\nCV validation failed:", message)
        print(json.dumps(data, indent=2, ensure_ascii=False))
        raise ValueError(message)

    return data


def as_list(value):
    if value is None:
        return []
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    if isinstance(value, str) and value.strip():
        return [value.strip()]
    return []


def join_items(value, separator=" | "):
    return separator.join(as_list(value))


def education_line(education, prefix):
    degree = str(education.get(f"{prefix}_degree", "") or "").strip()
    institution = str(education.get(f"{prefix}_institution", "") or "").strip()
    duration = str(education.get(f"{prefix}_duration", "") or "").strip()
    score = str(education.get(f"{prefix}_score", "") or "").strip()

    parts = [part for part in (degree, institution) if part]
    if score:
        label = "CGPA:" if "cgpa" not in score.lower() and "percentage" not in score.lower() else ""
        parts.append(f"{label} {score}".strip())
    if duration:
        parts.append(duration)
    return " | ".join(parts)


def build_replacements(cv):
    replacements = {
        "{{NAME}}": str(cv.get("name", "") or ""),
        "{{TARGETED_HEADLINE}}": str(
            cv.get("targeted_headline", "") or ""
        ),
        "{{LOCATION}}": str(cv.get("location", "") or ""),
        "{{PHONE}}": str(cv.get("phone", "") or ""),
        "{{EMAIL}}": str(cv.get("email", "") or ""),
        "{{LINKEDIN}}": str(cv.get("linkedin", "") or ""),
        "{{TARGETED_SUMMARY}}": str(cv.get("targeted_summary", "") or ""),
        "{{CORE_EXPERTISE}}": join_items(cv.get("core_expertise", [])),
        "{{TOOLS_AND_TECHNOLOGIES}}": join_items(
            cv.get("tools_and_technologies", []), separator="\n"
        ),
        "{{CERTIFICATIONS}}": join_items(cv.get("certifications", []), separator=" • "),
    }

    education = cv.get("education", {})
    if not isinstance(education, dict):
        education = {}
    replacements["{{EDUCATION_1}}"] = education_line(education, "mba")
    replacements["{{EDUCATION_2}}"] = education_line(education, "btech")

    experiences = cv.get("experience", [])
    if not isinstance(experiences, list):
        experiences = []

    for i in range(1, 5):
        exp = experiences[i - 1] if i <= len(experiences) and isinstance(experiences[i - 1], dict) else {}
        heading_parts = [
            str(exp.get("company", "") or "").strip(),
            str(exp.get("role", "") or "").strip(),
            str(exp.get("dates", "") or "").strip(),
        ]
        replacements[f"{{{{EXPERIENCE_{i}_HEADING}}}}"] = " | ".join(
            part for part in heading_parts if part
        )
        replacements[f"{{{{EXPERIENCE_{i}_FOCUS}}}}"] = (
            "Focus: " + str(exp.get("focus", "")).strip()
            if exp.get("focus") and i < 4 else ""
        )
        bullets = as_list(exp.get("bullets", []))
        bullet_count = max(EXPERIENCE_BULLET_LIMITS[i], len(bullets))
        for b in range(1, bullet_count + 1):
            replacements["{{EXPERIENCE_" + str(i) + "_BULLET_" + str(b) + "}}"] = (
                bullets[b - 1] if b <= len(bullets) else ""
            )

    achievements = as_list(cv.get("achievements", []))
    leadership = as_list(cv.get("leadership", []))
    combined = (achievements + leadership)[:5]
    for i in range(1, 6):
        replacements["{{ACHIEVEMENT_" + str(i) + "}}"] = (
            combined[i - 1] if i <= len(combined) else ""
        )

    return replacements


def replace_in_paragraph(paragraph, replacements):
    original = paragraph.text
    updated = original
    for placeholder, value in replacements.items():
        updated = updated.replace(placeholder, str(value))

    if updated == original:
        return

    # Keep the paragraph's first-run formatting instead of resetting paragraph.text.
    if paragraph.runs:
        paragraph.runs[0].text = updated
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(updated)


def replace_in_document(doc, replacements):
    for paragraph in doc.paragraphs:
        replace_in_paragraph(paragraph, replacements)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_in_paragraph(paragraph, replacements)


def expand_experience_bullet_slots(doc, cv):
    """Clone existing template bullet paragraphs when verified content exceeds its slots."""
    experiences = cv.get("experience", [])
    if not isinstance(experiences, list):
        return
    for role_index in range(1, min(4, len(experiences)) + 1):
        experience = experiences[role_index - 1]
        if not isinstance(experience, dict):
            continue
        bullets = as_list(experience.get("bullets", []))
        capacity = EXPERIENCE_BULLET_LIMITS[role_index]
        if len(bullets) <= capacity:
            continue
        marker = f"{{{{EXPERIENCE_{role_index}_BULLET_{capacity}}}}}"
        anchor = next((p for p in doc.paragraphs if marker in p.text), None)
        if anchor is None:
            raise ValueError(f"Template is missing the role {role_index} bullet anchor.")
        for bullet_index in range(capacity + 1, len(bullets) + 1):
            clone = deepcopy(anchor._p)
            new_marker = f"{{{{EXPERIENCE_{role_index}_BULLET_{bullet_index}}}}}"
            cloned_paragraph = Paragraph(clone, anchor._parent)
            replace_in_paragraph(cloned_paragraph, {marker: new_marker})
            anchor._p.addnext(clone)
            anchor = cloned_paragraph
            marker = new_marker


def _candidate_emphasis_spans(text):
    candidates = []
    leading = re.match(r"\s*(?:Successfully\s+)?([A-Za-z-]+)", text)
    if leading and leading.group(1).casefold() in ACTION_VERBS:
        candidates.append((leading.start(1), leading.end(1), 70))

    for pattern, priority in ((MONEY_PATTERN, 110), (PERCENT_PATTERN, 105), (SCALE_PATTERN, 100)):
        candidates.extend((match.start(), match.end(), priority) for match in pattern.finditer(text))

    folded_text = text.casefold()
    for phrase in ACHIEVEMENT_PHRASES:
        start = 0
        while True:
            index = folded_text.find(phrase.casefold(), start)
            if index < 0:
                break
            candidates.append((index, index + len(phrase), 60))
            start = index + len(phrase)

    selected = []
    for start, end, priority in sorted(candidates, key=lambda item: (-item[2], item[0], -(item[1] - item[0]))):
        if any(start < selected_end and end > selected_start for selected_start, selected_end in selected):
            continue
        selected.append((start, end))
        if len(selected) >= 3:
            break
    return sorted(selected)


def _paragraph_has_hyperlink(paragraph):
    return bool(paragraph._p.xpath(".//w:hyperlink"))


def apply_smart_bold_formatting(paragraph):
    """Bold a few deterministic spans while preserving each source run's properties."""
    if _paragraph_has_hyperlink(paragraph):
        return
    if not paragraph.style or paragraph.style.name != "List Bullet":
        return
    runs = list(paragraph.runs)
    text = "".join(run.text for run in runs)
    spans = _candidate_emphasis_spans(text)
    if not spans:
        return

    boundaries = {0, len(text)}
    offset = 0
    for run in runs:
        offset += len(run.text)
        boundaries.add(offset)
    already_partitioned = all(start in boundaries and end in boundaries for start, end in spans)
    if already_partitioned:
        offset = 0
        for run in runs:
            end = offset + len(run.text)
            if any(span_start <= offset and end <= span_end and end > offset for span_start, span_end in spans):
                run.bold = True
            offset = end
        return

    offset = 0
    for run in runs:
        run_text = run.text
        run_start, run_end = offset, offset + len(run_text)
        offset = run_end
        if not run_text or not any(start < run_end and end > run_start for start, end in spans):
            continue
        local_boundaries = {0, len(run_text)}
        for start, end in spans:
            if start < run_end and end > run_start:
                local_boundaries.add(max(0, start - run_start))
                local_boundaries.add(min(len(run_text), end - run_start))
        points = sorted(local_boundaries)
        clones = []
        for left, right in zip(points, points[1:]):
            if left == right:
                continue
            clone = deepcopy(run._r)
            cloned_run = Run(clone, paragraph)
            cloned_run.text = run_text[left:right]
            absolute_left, absolute_right = run_start + left, run_start + right
            if any(start <= absolute_left and absolute_right <= end for start, end in spans):
                cloned_run.bold = True
            clones.append(clone)
        parent = run._r.getparent()
        index = parent.index(run._r)
        parent.remove(run._r)
        for clone_index, clone in enumerate(clones):
            parent.insert(index + clone_index, clone)


def apply_smart_bold_to_document(doc):
    for paragraph in doc.paragraphs:
        apply_smart_bold_formatting(paragraph)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    apply_smart_bold_formatting(paragraph)


def find_unresolved_placeholders(doc):
    unresolved = []
    for paragraph in doc.paragraphs:
        if "{{" in paragraph.text or "}}" in paragraph.text:
            unresolved.append(paragraph.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if "{{" in paragraph.text or "}}" in paragraph.text:
                        unresolved.append(paragraph.text)
    return unresolved


def remove_empty_paragraphs(container):
    """Remove empty bullet/focus placeholders so no stray bullets or blank slots remain."""
    for paragraph in list(container.paragraphs):
        if paragraph.text.strip():
            continue
        style_name = paragraph.style.name.lower() if paragraph.style else ""
        if "bullet" in style_name or style_name == "normal":
            paragraph._element.getparent().remove(paragraph._element)

    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                remove_empty_paragraphs(cell)


def render_cv_document(cv_data, jd_data, output_dir=OUTPUT_DIR):
    if not TEMPLATE_FILE.is_file():
        raise FileNotFoundError(f"Required template not found: {TEMPLATE_FILE}")
    doc = Document(TEMPLATE_FILE)
    expand_experience_bullet_slots(doc, cv_data)
    replace_in_document(doc, build_replacements(cv_data))
    apply_smart_bold_to_document(doc)
    unresolved = find_unresolved_placeholders(doc)
    if unresolved:
        raise ValueError(
            "Unresolved template placeholders remain: " + " | ".join(unresolved[:10])
        )
    remove_empty_paragraphs(doc)

    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    job_title = str(jd_data.get("job_title", "Tailored_CV") or "Tailored_CV")
    safe_title = re.sub(r"[^A-Za-z0-9_-]+", "_", job_title).strip("_") or "Tailored_CV"
    candidate = str(cv_data.get("name", "Candidate") or "Candidate")
    safe_name = re.sub(r"[^A-Za-z0-9_-]+", "_", candidate).strip("_") or "Candidate"
    output_path = output_dir / f"{safe_name}_{safe_title}_CV.docx"
    doc.save(output_path)
    return output_path


def main():
    print("=" * 65)
    print("TAILORED CV GENERATOR")
    print("=" * 65)

    for path in (RESUME_FILE, JD_FILE, MATCH_FILE, TEMPLATE_FILE):
        if not path.exists():
            raise FileNotFoundError(f"Required file not found: {path}")

    print("\nLoading resume evidence...")
    resume_data = load_json(RESUME_FILE)
    print("Loading JD analysis...")
    jd_data = load_json(JD_FILE)
    print("Loading resume-JD matching...")
    match_data = load_json(MATCH_FILE)

    print("\nGenerating tailored CV using Gemini...")
    cv_data = generate_cv(resume_data, jd_data, match_data)
    # Always use verified certifications from the master-resume evidence.
    cv_data["certifications"] = normalize_source_certifications(resume_data.get("certifications", []))
    print("Verified master-resume certifications:", len(cv_data["certifications"]))

    GENERATED_JSON_FILE.parent.mkdir(parents=True, exist_ok=True)
    with open(GENERATED_JSON_FILE, "w", encoding="utf-8") as f:
        json.dump(cv_data, f, indent=2, ensure_ascii=False)
    print("Saved generated CV JSON:", GENERATED_JSON_FILE)

    print("Loading master-style Word template...")
    output_path = render_cv_document(cv_data, jd_data)

    print("\n" + "=" * 65)
    print("SUCCESS")
    print("=" * 65)
    print("CV saved to:", output_path)
    print("Generated JSON:", GENERATED_JSON_FILE)
    print("Template headings and order are preserved from the optimized master resume.")


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print("\nERROR:")
        print(str(exc))
        print("\nThe CV was not confirmed as successful.")
        raise
