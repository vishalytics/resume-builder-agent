import re
from pathlib import Path

from docx import Document


def _normalize(value):
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").lower()).strip()


def _source_evidence(resume_data):
    evidence = []
    summary = resume_data.get("summary")
    if isinstance(summary, str) and summary.strip():
        evidence.append(("Career summary", summary.strip()))

    for experience in resume_data.get("experience", []):
        if not isinstance(experience, dict):
            continue
        company = str(experience.get("company", "") or "").strip()
        role = str(experience.get("role", "") or "").strip()
        heading = " | ".join(value for value in (company, role) if value)
        for achievement in experience.get("achievements", []):
            if isinstance(achievement, dict):
                text = achievement.get("evidence") or achievement.get("achievement")
                metrics = achievement.get("metrics", [])
                if isinstance(metrics, list):
                    text = ". ".join([str(text or "")] + [str(item) for item in metrics if item])
            else:
                text = achievement
            if isinstance(text, str) and text.strip():
                evidence.append((heading or "Work experience", text.strip()))
        responsibilities = experience.get("responsibilities", [])
        if isinstance(responsibilities, list) and responsibilities:
            evidence.append((heading or "Work experience", "; ".join(map(str, responsibilities))))

    for item in resume_data.get("education", []):
        if isinstance(item, dict):
            line = " | ".join(
                str(item.get(key, "") or "").strip()
                for key in ("degree", "institution", "cgpa", "duration")
                if item.get(key)
            )
            if line:
                evidence.append(("Education", line))
    return evidence


def _find_evidence(evidence, patterns, limit=2):
    matches = []
    for heading, text in evidence:
        normalized = _normalize(text)
        if any(_normalize(pattern) in normalized for pattern in patterns):
            matches.append((heading, text))
        if len(matches) == limit:
            break
    return matches


def build_interview_questions(resume_data, jd_data, match_data):
    role = str(jd_data.get("job_title") or "the role")
    evidence = _source_evidence(resume_data)
    questions = [
        (
            f"What makes your experience a strong fit for {role}?",
            ("strategy", "founder", "cross-functional", "operations"),
        ),
        (
            "Describe a process you automated or improved. What changed as a result?",
            ("vendor-slot", "manual coordination", "automation", "workflow"),
        ),
        (
            "How have you used dashboards or operational data to support a decision?",
            ("dashboard", "power bi", "analytics", "data-driven"),
        ),
        (
            "Tell me about a cross-functional initiative you took from planning through rollout.",
            ("cross-functional", "implementation", "rollout", "vendor coordination"),
        ),
        (
            "Share an example of improving cost, efficiency, or execution at scale.",
            ("reduced costs", "20%", "operational efficiency", "600+", "1,200"),
        ),
        (
            "How have you supported market research, pricing, or go-to-market work?",
            ("market research", "competitor", "pricing", "gtm", "vendor onboarding"),
        ),
    ]

    no_evidence = match_data.get("no_evidence", [])
    if no_evidence:
        requirement = no_evidence[0].get("requirement", "a requirement")
        questions.append((
            f"How would you approach {requirement} when your direct experience is limited?",
            (),
        ))

    result = []
    for question, patterns in questions:
        matches = _find_evidence(evidence, patterns) if patterns else []
        if matches:
            points = [f"{heading}: {text}" for heading, text in matches]
            points.append("Structure the response as situation, action, and result; keep the cited figures unchanged.")
        else:
            points = [
                "The supplied resume evidence does not document this directly. Be explicit about that gap.",
                "Describe only the closest demonstrated experience and explain a practical approach without claiming prior ownership.",
            ]
        result.append({"question": question, "answer_points": points})
    return result


def generate_interview_document(resume_data, jd_data, match_data, output_path):
    document = Document()
    document.add_heading("Interview Preparation", 0)
    role = str(jd_data.get("job_title") or "Target role")
    company = str(jd_data.get("company") or "")
    document.add_paragraph(" | ".join(value for value in (role, company) if value))
    document.add_paragraph(
        "Suggested points below are drawn from the supplied resume evidence. "
        "Keep examples and figures accurate; acknowledge gaps directly."
    )

    for index, item in enumerate(build_interview_questions(resume_data, jd_data, match_data), start=1):
        document.add_heading(f"{index}. {item['question']}", level=1)
        document.add_paragraph("Suggested answer points", style="Heading 2")
        for point in item["answer_points"]:
            document.add_paragraph(point, style="List Bullet")

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path