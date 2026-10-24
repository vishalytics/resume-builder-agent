import json
import re
from pathlib import Path

from docx import Document
from gemini_client import generate_with_fallback


# ============================================================
# CONFIG
# ============================================================

RESUME_FILE = Path("data/resume_evidence.json")
JD_FILE = Path("data/jd_analysis.json")
MATCH_FILE = Path("data/resume_jd_match.json")

# This template is based on the optimized master resume layout.
TEMPLATE_FILE = Path("templates/CV_TEMPLATE.docx")

OUTPUT_DIR = Path("output")
GENERATED_JSON_FILE = Path("data/generated_cv.json")


def load_json(path):
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def compact_resume_data(data):
    """Keep source evidence needed for a tailored CV, including nested identity."""
    keys = [
        "candidate", "name", "email", "phone", "linkedin", "location",
        "summary", "education", "experience", "work_experience",
        "professional_experience", "employment", "skills", "technical_skills",
        "tools_and_technologies", "achievements", "certifications",
        "projects", "leadership",
    ]
    return {key: data[key] for key in keys if key in data}


def compact_jd_data(data):
    keys = [
        "job_title", "company", "domain_keywords", "ats_keywords",
        "behavioral_requirements", "experience_requirements",
        "education_requirements", "important_action_verbs",
        "key_business_outcomes",
    ]
    return {key: data[key] for key in keys if key in data}


def compact_match_data(data):
    keys = [
        "matches", "experience_summary", "education_summary",
        "tools_and_technologies", "key_strengths", "recommendations",
    ]
    return {key: data[key] for key in keys if key in data}


def validate_cv(data):
    if not isinstance(data, dict):
        return False, "Gemini response is not a JSON object."

    required = [
        "name", "targeted_summary", "core_expertise",
        "tools_and_technologies", "experience", "education",
        "achievements", "certifications",
    ]
    for field in required:
        if field not in data:
            return False, f"Missing field: {field}"

    if not str(data.get("name", "")).strip():
        return False, "Name is empty."
    if not str(data.get("targeted_summary", "")).strip():
        return False, "Targeted summary is empty."
    if not isinstance(data.get("experience"), list):
        return False, "Experience must be a list."
    if not isinstance(data.get("education"), dict):
        return False, "Education must be an object."

    for i, exp in enumerate(data["experience"], start=1):
        if not isinstance(exp, dict):
            return False, f"Experience entry {i} is not an object."
        if not isinstance(exp.get("bullets", []), list):
            return False, f"Bullets for experience entry {i} must be a list."

    return True, "Valid CV JSON."



def normalize_source_certifications(value):
    """Normalize certification entries from resume_evidence.json, never from model output."""
    if value is None:
        return []

    raw_items = value if isinstance(value, list) else [value]
    entries = []
    for item in raw_items:
        if isinstance(item, dict):
            item = (item.get("name") or item.get("certification") or
                    item.get("title") or item.get("credential") or "")
        if not isinstance(item, str):
            continue
        # The parser may preserve the whole CERTIFICATIONS line as one string.
        for part in re.split(r"\s*[•●▪]\s*|\r?\n+", item):
            part = re.sub(r"^\s*[-*]+\s*", "", part).strip()
            if not part:
                continue
            # Guard against the exact failure seen: education copied into certifications.
            education_markers = (
                "mba degree", "b.tech degree", "btech degree",
                "indian institute of technology", "national institute of technology",
                "cgpa:", "2023–2025", "2023-2025", "2014–2018", "2014-2018"
            )
            if any(marker in part.lower() for marker in education_markers):
                continue
            if part not in entries:
                entries.append(part)
    return entries

def generate_cv(resume_data, jd_data, match_data):
    resume = compact_resume_data(resume_data)
    jd = compact_jd_data(jd_data)
    match = compact_match_data(match_data)

    prompt = f"""
You are an expert ATS resume editor. Produce CV CONTENT ONLY; the Word template controls
the layout, typography, heading names, and section order.

TARGET ROLE:
{jd.get("job_title", "Target Role")}

NON-NEGOTIABLE RULES:
- Use only facts supported by CANDIDATE EVIDENCE below. Do not invent employers, titles,
  dates, qualifications, tools, metrics, achievements, or certifications.
- Keep the experience entries in the same order as the source resume.
- Tailor wording to the target role only where the source evidence supports it.
- Never include editorial notes such as "confirm", "validate", "add your exact role",
  "reported notes cite", or similar instructions in the final CV.
- Do not add or rename resume section headings. The Word template supplies the headings:
  EDUCATION; HOW I CREATE BUSINESS IMPACT; CORE COMPETENCIES; TOOLS & TECHNOLOGY;
  PROFESSIONAL EXPERIENCE; SELECTED ACHIEVEMENTS & LEADERSHIP; CERTIFICATIONS.
- Keep the summary concise (2–3 lines). Prefer measurable results when verified.
- Return valid JSON only, with no markdown fences or explanations.
- Use empty strings/lists for missing information. Never output placeholder braces.

JOB REQUIREMENTS:
{json.dumps(jd, ensure_ascii=False)}

RESUME-JD MATCH:
{json.dumps(match, ensure_ascii=False)}

CANDIDATE EVIDENCE:
{json.dumps(resume, ensure_ascii=False)}

Return exactly this JSON structure:
{{
  "name": "",
  "email": "",
  "phone": "",
  "linkedin": "",
  "location": "",
  "targeted_headline": "",
  "targeted_summary": "",
  "core_expertise": [],
  "tools_and_technologies": [],
  "experience": [
    {{
      "company": "",
      "role": "",
      "dates": "",
      "focus": "",
      "bullets": []
    }}
  ],
  "achievements": [],
  "leadership": [],
  "certifications": [],
  "education": {{
    "mba_degree": "",
    "mba_institution": "",
    "mba_duration": "",
    "mba_score": "",
    "btech_degree": "",
    "btech_institution": "",
    "btech_duration": "",
    "btech_score": ""
  }}
}}

Additional requirements:
- Return up to 4 experience entries in source order.
- Use up to 8 concise bullets for the first two roles, up to 5 for the third, and up to 2 for the fourth.
- "focus" should be a short line of verified themes, without inventing skills.
- "core_expertise" should contain concise competencies supported by the evidence.
- "tools_and_technologies" should be a list of short category lines, such as
  "Analytics: Excel, Power BI, SQL" and "Automation & AI: Google Apps Script, Gemini APIs".
- Put leadership roles in "leadership"; do not duplicate them in "achievements".
- Preserve degree, institution, dates, and scores as written in the evidence.
- Do not include projects as a separate section because the master layout has no projects heading.
"""


    print("\nSending request to Gemini...")
    result = generate_with_fallback(prompt)
    print("Gemini response received. Response length:", len(result))

    try:
        data = json.loads(result)
    except json.JSONDecodeError as exc:
        print("\nGemini returned invalid JSON:")
        print(result)
        raise ValueError("Invalid JSON returned by Gemini.") from exc

    # Identity fields are taken from trusted evidence when available, not invented by the model.
    candidate = resume_data.get("candidate", {})
    if not isinstance(candidate, dict):
        candidate = {}

    aliases = {
        "name": ("name", "full_name", "candidate_name"),
        "email": ("email", "email_address"),
        "phone": ("phone", "mobile", "phone_number"),
        "linkedin": ("linkedin", "linkedin_url", "linkedin_profile"),
        "location": ("location", "city", "current_location"),
    }
    for field, candidates in aliases.items():
        source_value = ""
        for key in candidates:
            if resume_data.get(key):
                source_value = resume_data[key]
                break
            if candidate.get(key):
                source_value = candidate[key]
                break
        if source_value:
            data[field] = str(source_value)

    # Certifications must come from the master-resume evidence, not Gemini's guess.
    # This prevents education degrees/institutions from being copied into CERTIFICATIONS.
    source_certs = normalize_source_certifications(resume_data.get("certifications", []))
    if source_certs:
        data["certifications"] = source_certs
    else:
        data["certifications"] = []
        print(
            "WARNING: No certification entries were found in data/resume_evidence.json. "
            "The Certifications section will be left blank rather than populated with education."
        )

    valid, message = validate_cv(data)
    if not valid:
        print("\nCV validation failed:", message)
        print(json.dumps(data, indent=2, ensure_ascii=False))
        raise ValueError(message)

    return data


def as_list(value):
    if value is None:
        return []
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    if isinstance(value, str) and value.strip():
        return [value.strip()]
    return []


def join_items(value, separator=" | "):
    return separator.join(as_list(value))


def education_line(education, prefix):
    degree = str(education.get(f"{prefix}_degree", "") or "").strip()
    institution = str(education.get(f"{prefix}_institution", "") or "").strip()
    duration = str(education.get(f"{prefix}_duration", "") or "").strip()
    score = str(education.get(f"{prefix}_score", "") or "").strip()

    parts = [part for part in (degree, institution) if part]
    if score:
        label = "CGPA:" if "cgpa" not in score.lower() and "percentage" not in score.lower() else ""
        parts.append(f"{label} {score}".strip())
    if duration:
        parts.append(duration)
    return " | ".join(parts)


def build_replacements(cv):
    replacements = {
        "{{NAME}}": str(cv.get("name", "") or ""),
        "{{TARGETED_HEADLINE}}": str(
            cv.get("targeted_headline", "") or
            "STRATEGY & BUSINESS OPERATIONS | FOUNDER’S OFFICE | AI AUTOMATION & ANALYTICS"
        ),
        "{{LOCATION}}": str(cv.get("location", "") or ""),
        "{{PHONE}}": str(cv.get("phone", "") or ""),
        "{{EMAIL}}": str(cv.get("email", "") or ""),
        "{{LINKEDIN}}": str(cv.get("linkedin", "") or ""),
        "{{TARGETED_SUMMARY}}": str(cv.get("targeted_summary", "") or ""),
        "{{CORE_EXPERTISE}}": join_items(cv.get("core_expertise", [])),
        "{{TOOLS_AND_TECHNOLOGIES}}": join_items(
            cv.get("tools_and_technologies", []), separator="\n"
        ),
        "{{CERTIFICATIONS}}": join_items(cv.get("certifications", []), separator=" • "),
    }

    education = cv.get("education", {})
    if not isinstance(education, dict):
        education = {}
    replacements["{{EDUCATION_1}}"] = education_line(education, "mba")
    replacements["{{EDUCATION_2}}"] = education_line(education, "btech")

    experiences = cv.get("experience", [])
    if not isinstance(experiences, list):
        experiences = []

    max_bullets = {1: 8, 2: 8, 3: 5, 4: 2}
    for i in range(1, 5):
        exp = experiences[i - 1] if i <= len(experiences) and isinstance(experiences[i - 1], dict) else {}
        heading_parts = [
            str(exp.get("company", "") or "").strip(),
            str(exp.get("role", "") or "").strip(),
            str(exp.get("dates", "") or "").strip(),
        ]
        replacements[f"{{{{EXPERIENCE_{i}_HEADING}}}}"] = " | ".join(
            part for part in heading_parts if part
        )
        replacements[f"{{{{EXPERIENCE_{i}_FOCUS}}}}"] = (
            "Focus: " + str(exp.get("focus", "")).strip()
            if exp.get("focus") and i < 4 else ""
        )
        bullets = as_list(exp.get("bullets", []))
        for b in range(1, max_bullets[i] + 1):
            replacements["{{EXPERIENCE_" + str(i) + "_BULLET_" + str(b) + "}}"] = (
                bullets[b - 1] if b <= len(bullets) else ""
            )

    achievements = as_list(cv.get("achievements", []))
    leadership = as_list(cv.get("leadership", []))
    combined = (achievements + leadership)[:5]
    for i in range(1, 6):
        replacements["{{ACHIEVEMENT_" + str(i) + "}}"] = (
            combined[i - 1] if i <= len(combined) else ""
        )

    return replacements


def replace_in_paragraph(paragraph, replacements):
    original = paragraph.text
    updated = original
    for placeholder, value in replacements.items():
        updated = updated.replace(placeholder, str(value))

    if updated == original:
        return

    # Keep the paragraph's first-run formatting instead of resetting paragraph.text.
    if paragraph.runs:
        paragraph.runs[0].text = updated
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(updated)


def replace_in_document(doc, replacements):
    for paragraph in doc.paragraphs:
        replace_in_paragraph(paragraph, replacements)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_in_paragraph(paragraph, replacements)


def find_unresolved_placeholders(doc):
    unresolved = []
    for paragraph in doc.paragraphs:
        if "{{" in paragraph.text or "}}" in paragraph.text:
            unresolved.append(paragraph.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if "{{" in paragraph.text or "}}" in paragraph.text:
                        unresolved.append(paragraph.text)
    return unresolved


def remove_empty_paragraphs(container):
    """Remove empty bullet/focus placeholders so no stray bullets or blank slots remain."""
    for paragraph in list(container.paragraphs):
        if paragraph.text.strip():
            continue
        style_name = paragraph.style.name.lower() if paragraph.style else ""
        if "bullet" in style_name or style_name == "normal":
            paragraph._element.getparent().remove(paragraph._element)

    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                remove_empty_paragraphs(cell)


def main():
    print("=" * 65)
    print("TAILORED CV GENERATOR")
    print("=" * 65)

    for path in (RESUME_FILE, JD_FILE, MATCH_FILE, TEMPLATE_FILE):
        if not path.exists():
            raise FileNotFoundError(f"Required file not found: {path}")

    print("\nLoading resume evidence...")
    resume_data = load_json(RESUME_FILE)
    print("Loading JD analysis...")
    jd_data = load_json(JD_FILE)
    print("Loading resume-JD matching...")
    match_data = load_json(MATCH_FILE)

    print("\nGenerating tailored CV using Gemini...")
    cv_data = generate_cv(resume_data, jd_data, match_data)
    # Always use verified certifications from the master-resume evidence.
    cv_data["certifications"] = normalize_source_certifications(resume_data.get("certifications", []))
    print("Verified master-resume certifications:", len(cv_data["certifications"]))

    GENERATED_JSON_FILE.parent.mkdir(parents=True, exist_ok=True)
    with open(GENERATED_JSON_FILE, "w", encoding="utf-8") as f:
        json.dump(cv_data, f, indent=2, ensure_ascii=False)
    print("Saved generated CV JSON:", GENERATED_JSON_FILE)

    print("Loading master-style Word template...")
    doc = Document(TEMPLATE_FILE)
    replace_in_document(doc, build_replacements(cv_data))
    unresolved = find_unresolved_placeholders(doc)
    if unresolved:
        raise ValueError(
            "Unresolved template placeholders remain: " + " | ".join(unresolved[:10])
        )
    remove_empty_paragraphs(doc)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    job_title = str(jd_data.get("job_title", "Tailored_CV") or "Tailored_CV")
    safe_title = re.sub(r"[^A-Za-z0-9_-]+", "_", job_title).strip("_") or "Tailored_CV"
    output_path = OUTPUT_DIR / f"Vishal_Kumar_{safe_title}_CV.docx"
    doc.save(output_path)

    print("\n" + "=" * 65)
    print("SUCCESS")
    print("=" * 65)
    print("CV saved to:", output_path)
    print("Generated JSON:", GENERATED_JSON_FILE)
    print("Template headings and order are preserved from the optimized master resume.")


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print("\nERROR:")
        print(str(exc))
        print("\nThe CV was not confirmed as successful.")
        raise
