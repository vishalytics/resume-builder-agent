import json
import re
from pathlib import Path

from docx import Document

from gemini_client import generate_with_fallback, parse_json_response


def _normalize(value):
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").lower()).strip()


def _source_evidence(resume_data):
    evidence = []
    summary = resume_data.get("summary")
    if isinstance(summary, str) and summary.strip():
        evidence.append(("Career summary", summary.strip()))
    for experience in resume_data.get("experience", []):
        if not isinstance(experience, dict):
            continue
        company = str(experience.get("company", "") or "").strip()
        role = str(experience.get("role", "") or "").strip()
        heading = " | ".join(value for value in (company, role) if value) or "Work experience"
        for achievement in experience.get("achievements", []):
            if isinstance(achievement, dict):
                text = achievement.get("evidence") or achievement.get("achievement")
                metrics = achievement.get("metrics", [])
                if isinstance(metrics, list):
                    text = ". ".join([str(text or "")] + [str(item) for item in metrics if item])
            else:
                text = achievement
            if isinstance(text, str) and text.strip():
                evidence.append((heading, text.strip()))
        responsibilities = experience.get("responsibilities", [])
        if isinstance(responsibilities, list) and responsibilities:
            evidence.append((heading, "; ".join(map(str, responsibilities))))
    for item in resume_data.get("education", []):
        if isinstance(item, dict):
            line = " | ".join(
                str(item.get(key, "") or "").strip()
                for key in ("degree", "institution", "cgpa", "duration")
                if item.get(key)
            )
            if line:
                evidence.append(("Education", line))
    return evidence


def _validate_interview_response(response_text, evidence_count, gap_count):
    try:
        result = parse_json_response(response_text)
    except ValueError:
        return False, "invalid_json"
    questions = result.get("questions")
    if not isinstance(questions, list) or not 4 <= len(questions) <= 10:
        return False, "invalid_question_count"
    categories = set()
    for item in questions:
        if not isinstance(item, dict) or not isinstance(item.get("question"), str) or not item["question"].strip():
            return False, "invalid_question"
        if item.get("category") not in {"role_specific", "behavioral", "analytical_case"}:
            return False, "invalid_question_category"
        categories.add(item["category"])
        evidence_indices = item.get("evidence_indices")
        followups = item.get("follow_up_questions")
        if not isinstance(evidence_indices, list) or any(
            type(index) is not int or not 0 <= index < evidence_count
            for index in evidence_indices
        ):
            return False, "invalid_evidence_reference"
        if not isinstance(followups, list) or not 2 <= len(followups) <= 3 or any(
            not isinstance(value, str) or not value.strip() for value in followups
        ):
            return False, "invalid_follow_up_questions"
    gap_indices = result.get("experience_gap_indices")
    if not isinstance(gap_indices, list) or any(
        type(index) is not int or not 0 <= index < gap_count for index in gap_indices
    ):
        return False, "invalid_gap_reference"
    if "analytical_case" not in categories or "role_specific" not in categories:
        return False, "missing_question_category"
    if set(gap_indices) != set(range(gap_count)):
        return False, "incomplete_experience_gaps"
    return True, "valid"


def build_interview_questions(resume_data, jd_data, match_data, event_sink=None):
    role = str(jd_data.get("job_title") or "the role")
    evidence = _source_evidence(resume_data)
    gaps = match_data.get("no_evidence", [])
    if not isinstance(gaps, list):
        gaps = []
    jd_requirements = {
        key: jd_data.get(key, [])
        for key in ("responsibilities", "required_skills", "preferred_skills", "key_business_outcomes")
    }
    prompt = f"""
Create an interview question plan for {role}. Use the JD and numbered, verified
resume evidence below. Return valid JSON only. Do not draft autobiographical
claims or invented answer content: answer points will be assembled from source
evidence by the application. Include 6-8 questions covering role-specific
responsibilities, behavioral examples, and at least one analytical/case question.
For each item provide a question, category (role_specific, behavioral, or
analytical_case), evidence_indices (indices from the evidence list only), and
2-3 follow_up_questions. Reference no evidence for an analytical case when none
is directly relevant. experience_gap_indices must refer only to the listed
no-evidence requirements. Use [] when there are no documented gaps.

JD REQUIREMENTS:
{json.dumps(jd_requirements, ensure_ascii=False)}

VERIFIED RESUME EVIDENCE:
{json.dumps([{"index": index, "source": source, "text": text} for index, (source, text) in enumerate(evidence)], ensure_ascii=False)}

VERIFIED EXPERIENCE GAPS:
{json.dumps(gaps, ensure_ascii=False)}

Return exactly {{"questions":[{{"question":"", "category":"role_specific", "evidence_indices":[], "follow_up_questions":[]}}], "experience_gap_indices":[]}}
"""

    def validator(response_text):
        return _validate_interview_response(response_text, len(evidence), len(gaps))

    plan = parse_json_response(generate_with_fallback(
        prompt,
        stage="interview_preparation",
        response_validator=validator,
        event_sink=event_sink,
    ))
    questions = []
    for item in plan["questions"]:
        references = [evidence[index] for index in item["evidence_indices"]]
        if references:
            points = [f"{source}: {text}" for source, text in references]
            points.append("Use a situation, action, result structure and keep all source figures unchanged.")
        elif item["category"] == "analytical_case":
            points = [
                "Clarify the objective, available data, constraints, and decision criteria.",
                "Explain a proposed analytical approach as a case response, not as prior work experience.",
            ]
        else:
            points = [
                "No directly relevant example was selected from the supplied evidence.",
                "Acknowledge the gap and describe adjacent experience without claiming unsupported ownership.",
            ]
        questions.append({
            "question": item["question"],
            "category": item["category"],
            "answer_points": points,
            "follow_up_questions": item["follow_up_questions"],
            "evidence_indices": item["evidence_indices"],
        })
    selected_gaps = [gaps[index] for index in plan["experience_gap_indices"]]
    return {"questions": questions, "experience_gaps": selected_gaps}


def render_interview_document(content, jd_data, output_path):
    document = Document()
    document.add_heading("Interview Preparation", 0)
    role = str(jd_data.get("job_title") or "Target role")
    company = str(jd_data.get("company") or "")
    document.add_paragraph(" | ".join(value for value in (role, company) if value))
    document.add_paragraph(
        "Suggested evidence points are quoted from the resume evidence. "
        "Analytical-case guidance is a proposed approach, not a claim of past work."
    )
    for index, item in enumerate(content["questions"], start=1):
        document.add_heading(f"{index}. {item['question']}", level=1)
        document.add_paragraph("Suggested answer points", style="Heading 2")
        for point in item["answer_points"]:
            document.add_paragraph(point, style="List Bullet")
        document.add_paragraph("Likely follow-up questions", style="Heading 2")
        for follow_up in item["follow_up_questions"]:
            document.add_paragraph(follow_up, style="List Bullet")
    document.add_heading("Genuine experience gaps", level=1)
    if content["experience_gaps"]:
        for gap in content["experience_gaps"]:
            requirement = str(gap.get("requirement", "Unverified requirement"))
            reason = str(gap.get("reason", "Direct evidence is not documented in the supplied resume."))
            document.add_paragraph(f"{requirement}: {reason}", style="List Bullet")
    else:
        document.add_paragraph("No explicit no-evidence requirements were supplied by the resume matcher.")
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path


def generate_interview_document(resume_data, jd_data, match_data, output_path, event_sink=None):
    content = build_interview_questions(resume_data, jd_data, match_data, event_sink)
    return content, render_interview_document(content, jd_data, output_path)