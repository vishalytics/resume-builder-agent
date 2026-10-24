import hashlib
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

import cv_generator


class CvSmartBoldTests(unittest.TestCase):
    def test_action_and_percentage_are_bold_without_bolding_filler(self):
        document = Document()
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run("Reduced manual team effort by 40% by automating recurring reporting workflows.")
        original = paragraph.text
        cv_generator.apply_smart_bold_formatting(paragraph)
        bold_text = [run.text for run in paragraph.runs if run.bold]
        self.assertEqual(paragraph.text, original)
        self.assertIn("Reduced", bold_text)
        self.assertIn("40%", bold_text)
        self.assertFalse(any("manual team effort" in text or "recurring reporting" in text for text in bold_text))
        self.assertLessEqual(len(bold_text), 3)

    def test_counts_money_hyphenated_scale_and_unicode_remain_intact(self):
        examples = (
            "Managed a 17-member team across 8 residential campuses.",
            "Built a performance MIS covering 1,200+ students.",
            "Closed corporate orders worth ₹8 lakh.",
        )
        document = Document()
        paragraphs = [document.add_paragraph(text, style="List Bullet") for text in examples]
        for paragraph in paragraphs:
            original = paragraph.text
            cv_generator.apply_smart_bold_formatting(paragraph)
            self.assertEqual(paragraph.text, original)
            self.assertTrue(any(run.bold for run in paragraph.runs))
        bold = [[run.text for run in paragraph.runs if run.bold] for paragraph in paragraphs]
        self.assertIn("17-member team", bold[0])
        self.assertIn("8 residential campuses", bold[0])
        self.assertIn("1,200+ students", bold[1])
        self.assertIn("₹8 lakh", bold[2])

    def test_run_formatting_multiline_and_repeated_application_are_preserved(self):
        document = Document()
        paragraph = document.add_paragraph(style="List Bullet")
        first = paragraph.add_run("Improved efficiency\nby ")
        first.font.name = "Arial"
        first.font.size = Pt(8.5)
        first.italic = True
        second = paragraph.add_run("50%.")
        second.font.name = "Arial"
        second.font.size = Pt(8.5)
        original_text = paragraph.text
        paragraph_style = paragraph.style.name
        cv_generator.apply_smart_bold_formatting(paragraph)
        snapshot = [(run.text, run.bold, run.italic, run.font.name, run.font.size.pt if run.font.size else None) for run in paragraph.runs]
        self.assertEqual(paragraph.text, original_text)
        self.assertEqual(paragraph.style.name, paragraph_style)
        self.assertTrue(any(text == "Improved" and bold for text, bold, *_ in snapshot))
        self.assertTrue(any(text == "50%" and bold for text, bold, *_ in snapshot))
        self.assertTrue(all(font == "Arial" and size == 8.5 for _, _, _, font, size in snapshot))
        first_source_run = {"Improved", " efficiency\nby "}
        self.assertTrue(all(
            italic is True if text in first_source_run else italic is None
            for text, _, italic, _, _ in snapshot
        ))
        cv_generator.apply_smart_bold_formatting(paragraph)
        self.assertEqual(
            [(run.text, run.bold, run.italic, run.font.name, run.font.size.pt if run.font.size else None) for run in paragraph.runs],
            snapshot,
        )

    def test_dates_generic_sentences_and_non_bullet_paragraphs_are_not_bolded(self):
        document = Document()
        date = document.add_paragraph("Jul 2025 - Present", style="List Bullet")
        generic = document.add_paragraph("The role involved work across the business.", style="List Bullet")
        heading = document.add_paragraph("Built a report", style="Heading 2")
        for paragraph in (date, generic, heading):
            before = paragraph.text
            cv_generator.apply_smart_bold_formatting(paragraph)
            self.assertEqual(paragraph.text, before)
            self.assertFalse(any(run.bold for run in paragraph.runs))

    def test_hyperlink_paragraph_is_left_untouched(self):
        document = Document()
        paragraph = document.add_paragraph(style="List Bullet")
        relationship = paragraph.part.relate_to("https://example.invalid", RT.HYPERLINK, is_external=True)
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), relationship)
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "Reduced by 40%"
        run.append(text)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)
        before = paragraph._p.xml
        cv_generator.apply_smart_bold_formatting(paragraph)
        self.assertEqual(paragraph._p.xml, before)

    def test_expanded_verified_bullets_keep_template_style_and_section_order(self):
        cv = {
            "name": "Candidate", "targeted_headline": "Analyst", "targeted_summary": "Summary",
            "core_expertise": [], "tools_and_technologies": [], "education": {},
            "experience": [{
                "company": "Example Co", "role": "Analyst", "dates": "2020-2024",
                "bullets": [f"Built a verified dashboard for {index + 1},200+ students." for index in range(30)],
            }],
            "achievements": [], "leadership": [], "certifications": [],
        }
        template = Document(cv_generator.TEMPLATE_FILE)
        template_hash = hashlib.sha256(Path(cv_generator.TEMPLATE_FILE).read_bytes()).hexdigest()
        baseline_style = next(p.style.name for p in template.paragraphs if "{{EXPERIENCE_1_BULLET_1}}" in p.text)
        source_bullet = next(p for p in template.paragraphs if "{{EXPERIENCE_1_BULLET_1}}" in p.text)

        def paragraph_layout(paragraph):
            formatting = paragraph.paragraph_format
            return (
                paragraph.style.name,
                paragraph.alignment,
                formatting.left_indent,
                formatting.right_indent,
                formatting.first_line_indent,
                formatting.space_before,
                formatting.space_after,
                formatting.line_spacing,
                formatting.keep_together,
                formatting.keep_with_next,
            )

        source_layout = paragraph_layout(source_bullet)
        with tempfile.TemporaryDirectory() as folder:
            path = cv_generator.render_cv_document(cv, {"job_title": "Analyst"}, folder)
            output = Document(path)
        text = "\n".join(paragraph.text for paragraph in output.paragraphs)
        bullets = [p for p in output.paragraphs if p.style.name == baseline_style and p.text.startswith("Built a verified dashboard")]
        self.assertEqual(len(bullets), 30)
        self.assertIn("EDUCATION", text)
        self.assertIn("PROFESSIONAL EXPERIENCE", text)
        self.assertIn("CERTIFICATIONS", text)
        self.assertLess(text.index("EDUCATION"), text.index("PROFESSIONAL EXPERIENCE"))
        self.assertFalse(any("{{" in p.text for p in output.paragraphs))
        template_headings = [p.text for p in template.paragraphs if p.style.name == "Heading 1"]
        output_headings = [p.text for p in output.paragraphs if p.style.name == "Heading 1"]
        self.assertEqual(output_headings, template_headings)
        for source_section, output_section in zip(template.sections, output.sections):
            self.assertEqual(source_section.top_margin, output_section.top_margin)
            self.assertEqual(source_section.bottom_margin, output_section.bottom_margin)
            self.assertEqual(source_section.left_margin, output_section.left_margin)
            self.assertEqual(source_section.right_margin, output_section.right_margin)
            self.assertEqual(source_section.page_width, output_section.page_width)
            self.assertEqual(source_section.page_height, output_section.page_height)
        source_name = next(p for p in template.paragraphs if p.style.name.startswith("Heading") and p.text == "EDUCATION")
        output_name = next(p for p in output.paragraphs if p.style.name.startswith("Heading") and p.text == "EDUCATION")
        self.assertEqual(source_name.style.name, output_name.style.name)
        self.assertEqual(source_name.runs[0].font.name, output_name.runs[0].font.name)
        self.assertEqual(source_name.runs[0].font.size, output_name.runs[0].font.size)
        self.assertEqual(source_name.runs[0].bold, output_name.runs[0].bold)
        output_bullet = next(p for p in bullets if p.text == cv["experience"][0]["bullets"][0])
        self.assertEqual(paragraph_layout(output_bullet), source_layout)
        self.assertEqual(
            hashlib.sha256(Path(cv_generator.TEMPLATE_FILE).read_bytes()).hexdigest(),
            template_hash,
        )


if __name__ == "__main__":
    unittest.main()