from docx import Document
from pathlib import Path

RESUME_PATH = Path("master_resume/master_resume.docx")


def extract_resume_text():
    if not RESUME_PATH.exists():
        raise FileNotFoundError(
            f"Resume not found at: {RESUME_PATH}"
        )

    document = Document(RESUME_PATH)

    text_parts = []

    # Extract normal paragraphs
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            text_parts.append(text)

    # Extract tables
    for table in document.tables:
        for row in table.rows:
            row_text = []

            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    row_text.append(cell_text)

            if row_text:
                text_parts.append(" | ".join(row_text))

    return "\n".join(text_parts)


if __name__ == "__main__":
    resume_text = extract_resume_text()

    print("=" * 60)
    print("MASTER RESUME EXTRACTED")
    print("=" * 60)
    print(resume_text)
    print("=" * 60)

    print(f"\nCharacters extracted: {len(resume_text)}")